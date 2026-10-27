"""Build a plain one-column handout .docx from Markdown (the presentation prep sheet).

    python tools/build_prep_docx.py 3_Sources/Final_Demo_and_Presentation_Prep.md \
        2_Presentation_Prep/Final_Demo_and_Presentation_Prep.docx

Unlike build_report_docx.py this uses no template: A4, Calibri, Word's built-in styles.

Markup:
  # Title   ## Section                 document title / section heading
  - text   - [ ] text                  bullet / checkbox
  1. text                              numbered step
     indented line                     continues the list item above it
  | a | b |  rows                      table (first row = header, |---| row skipped)
  **bold**  *italic*  `code`           inline
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")
CODE_FONT = "Consolas"
HEADER_FILL = "D9E2F3"
PAGE_WIDTH = Cm(17.0)  # A4 less 2 cm margins


def add_inline(p, text: str, bold: bool = False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            add_inline(p, part[2:-2], bold=True)
        elif part.startswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = CODE_FONT
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
            r.bold = bold
        elif part.startswith("*"):
            r = p.add_run(part[1:-1])
            r.italic = True
            r.bold = bold
        else:
            p.add_run(part).bold = bold


def shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def table(doc, rows: list[list[str]]):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, cells in enumerate(rows):
        for j, text in enumerate(cells):
            cell = t.cell(i, j)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, text, bold=(i == 0))
            for r in p.runs:
                r.font.size = Pt(9.5) if r.font.name != CODE_FONT else Pt(8.5)
            if i == 0:
                shade(cell, HEADER_FILL)
    # Column widths follow the longest text in each column (capped), so a "Step" column stays narrow.
    weights = [max(len(rows[0][j]) + 3, 8, min(60, max(len(r[j]) for r in rows)))
               for j in range(len(rows[0]))]
    t.autofit = False
    grid = t._tbl.tblGrid.findall(qn("w:gridCol"))
    for j, w in enumerate(weights):
        width = int(PAGE_WIDTH * w / sum(weights))
        grid[j].set(qn("w:w"), str(width // 635))  # EMU to twentieths of a point
        for cell in t.columns[j].cells:
            cell.width = width
    repeat_header(t.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def build(src: Path, out: Path):
    doc = docx.Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.0))
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)

    lines = src.read_text(encoding="utf-8").splitlines()
    i, last_item = 0, None
    while i < len(lines):
        line = lines[i]
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[\s\-:|]+\|$", lines[i].strip()):
                    rows.append(split_row(lines[i]))
                i += 1
            table(doc, rows)
            last_item = None
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:], level=0)
            last_item = None
        elif line.startswith("## "):
            h = doc.add_heading(line[3:], level=1)
            h.paragraph_format.keep_with_next = True
            last_item = None
        elif line.startswith("- [ ] "):
            last_item = doc.add_paragraph()
            last_item.paragraph_format.left_indent = Cm(1.0)
            last_item.paragraph_format.first_line_indent = Cm(-0.6)
            add_inline(last_item, "☐  " + line[6:])
        elif line.startswith("- "):
            last_item = doc.add_paragraph(style="List Bullet")
            add_inline(last_item, line[2:])
        elif re.match(r"^\d+\. ", line):
            last_item = doc.add_paragraph(style="List Number")
            add_inline(last_item, line.split(". ", 1)[1])
        elif line.startswith("   ") and last_item is not None:
            last_item.add_run().add_break()
            add_inline(last_item, line.strip())
        elif line.strip():
            add_inline(doc.add_paragraph(), line)
            last_item = None
        i += 1
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("src", type=Path)
    ap.add_argument("out", type=Path)
    a = ap.parse_args()
    build(a.src, a.out)

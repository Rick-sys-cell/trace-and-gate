"""Build an IEEE-style report .docx from Markdown, reusing the interim report's layout.

    python tools/build_report_docx.py 3_Sources/Final_Report.md 4_Final_Submission/Final_Report.docx \
        --kind "Final Report"

The template (1_Interim_Submission/Interim_Report.docx, itself built from the proposal's layout)
supplies the cover page, the single-column title block, the two-column body and every paragraph
format: each element of the Markdown is a deep copy of the matching paragraph in the template,
with new text. Nothing else about the template is changed.

Markup:
  TITLE: ... / ABSTRACT: ... / INDEX: ...      title block
  # I. Heading   ## A. Heading                   section / subsection
  - *Lead.* text                               bullet (lead in italics)
  TABLE I: Title  followed by  | a | b | rows   table (first row = header)
  ![Fig. 1.  Caption](path/to/figure.png)      column-width figure (path relative to the .md)
  $$eq1  $$eq2                                 the hypotheses' equations (1) and (2)
  [n] Reference text                           reference (under # References)
  *italic*  **bold**                           inline
  <!-- ... -->                                 comment, ignored
"""
from __future__ import annotations

import argparse
import copy
import re
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import Emu

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "1_Interim_Submission" / "Interim_Report.docx"
COL_WIDTH_DXA = 4896  # one IEEE column, as in the template's tables
FIG_WIDTH_EMU = 3108960  # 3.4 in, as in the template's figures


class Template:
    """Index the template body once; prototypes are found by their text, not by position."""

    def __init__(self, path: Path):
        self.doc = docx.Document(str(path))
        self.body = self.doc.element.body
        self.els = list(self.body.iterchildren())
        text = [self._text(e) for e in self.els]

        def find(pred, what):
            for el, t in zip(self.els, text):
                if pred(el, t):
                    return el
            raise SystemExit(f"template: no {what} paragraph found")

        is_p = lambda el: el.tag == qn("w:p")  # noqa: E731
        style = lambda el: el.xpath("string(./w:pPr/w:pStyle/@w:val)")  # noqa: E731
        self.abstract = find(lambda e, t: t.startswith("Abstract"), "abstract")
        self.index = find(lambda e, t: t.startswith("Index Terms"), "index terms")
        self.h1 = find(lambda e, t: is_p(e) and style(e) == "Heading1", "Heading 1")
        self.h2 = find(lambda e, t: is_p(e) and style(e) == "Heading2", "Heading 2")
        i_h1 = self.els.index(self.h1)
        self.body_p = self.els[i_h1 + 1]
        self.tab_num = find(lambda e, t: re.fullmatch(r"TABLE [IVX]+", t), "TABLE n")
        self.tab_title = self.els[self.els.index(self.tab_num) + 1]
        self.table = next(e for e in self.els if e.tag == qn("w:tbl") and self.els.index(e) > i_h1)
        self.fig = find(lambda e, t: is_p(e) and e.xpath(".//w:drawing") and self.els.index(e) > i_h1, "figure")
        self.fig_cap = find(lambda e, t: t.startswith("Fig. 1."), "figure caption")
        self.bullet = find(lambda e, t: t.startswith("• "), "bullet")
        self.eq = [find(lambda e, t: t.endswith("(1)"), "equation 1"), find(lambda e, t: t.endswith("(2)"), "equation 2")]
        self.ref = find(lambda e, t: t.startswith("[1]"), "reference")
        self.title_paras = [e for e, t in zip(self.els[:i_h1], text[:i_h1]) if t.startswith("Trace and Gate")]
        self.kind_paras = [e for e, t in zip(self.els[:i_h1], text[:i_h1]) if t == "Interim Report"]
        self.venue_paras = [e for e, t in zip(self.els[:i_h1], text[:i_h1]) if "Interim Report, National" in t]
        self.first_body = self.els[self.els.index(self.index) + 1]
        self.final_sect = self.els[-1]

    @staticmethod
    def _text(el):
        return "".join(el.xpath(".//w:t/text()")) if el.tag == qn("w:p") else ""


def first_rpr(p):
    r = p.find(qn("w:r"))
    rpr = r.find(qn("w:rPr")) if r is not None else None
    return copy.deepcopy(rpr) if rpr is not None else None


def set_flag(rpr, tag, on):
    el = rpr.find(qn(tag))
    if on and el is None:
        el = rpr.makeelement(qn(tag), {})
        rpr.insert(0, el)
    elif not on and el is not None:
        rpr.remove(el)


def make_run(p, text, base_rpr, italic=None, bold=None):
    r = p.makeelement(qn("w:r"), {})
    rpr = copy.deepcopy(base_rpr) if base_rpr is not None else r.makeelement(qn("w:rPr"), {})
    if italic is not None:
        set_flag(rpr, "w:i", italic)
    if bold is not None:
        for b in rpr.findall(qn("w:b")):
            rpr.remove(b)
        if bold:
            set_flag(rpr, "w:b", True)
    if len(rpr):
        r.append(rpr)
    parts = text.split("\t")
    for i, part in enumerate(parts):
        if i:
            r.append(r.makeelement(qn("w:tab"), {}))
        if part:
            t = r.makeelement(qn("w:t"), {qn("xml:space"): "preserve"})
            t.text = part
            r.append(t)
    p.append(r)


INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def typographic(text: str) -> str:
    """Straight quotes to curly ones, as Word would type them; `code` loses its backticks."""
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r'(^|[\s(\[—–])"', "\\1\u201c", text)
    text = text.replace('"', "\u201d")
    text = re.sub(r"(^|[\s(\[—–])'", "\\1\u2018", text)
    return text.replace("'", "\u2019")


def fill(proto, text, base_rpr=None, italic_default=False):
    """Copy `proto`, drop its runs, and write `text` with *italic* / **bold** spans."""
    text = typographic(text)
    p = copy.deepcopy(proto)
    rpr = base_rpr if base_rpr is not None else first_rpr(proto)
    for r in p.findall(qn("w:r")):
        p.remove(r)
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**"):
            make_run(p, chunk[2:-2], rpr, italic=italic_default, bold=True)
        elif chunk.startswith("*"):
            make_run(p, chunk[1:-1], rpr, italic=not italic_default)
        else:
            make_run(p, chunk, rpr, italic=italic_default)
    return p


def set_runs(p, texts):
    """Replace the text of the first len(texts) runs, keeping their formatting; drop the rest."""
    runs = p.findall(qn("w:r"))
    for r, t in zip(runs, texts):
        for x in r.findall(qn("w:t")) + r.findall(qn("w:tab")):
            r.remove(x)
        make_tmp = p.makeelement(qn("w:t"), {qn("xml:space"): "preserve"})
        make_tmp.text = t
        r.append(make_tmp)
    for r in runs[len(texts):]:
        p.remove(r)


def table(tpl, rows):
    """Booktabs-style table like the template's: header rule above and below, rule under the last row."""
    proto_tbl = tpl.table
    cell_p = proto_tbl.find(".//" + qn("w:p"))
    head_rpr, body_rpr = first_rpr(cell_p), first_rpr(cell_p)
    set_flag(head_rpr, "w:b", True)
    for b in body_rpr.findall(qn("w:b")):
        body_rpr.remove(b)
    off = body_rpr.makeelement(qn("w:b"), {qn("w:val"): "0"})
    body_rpr.insert(0, off)
    ncol = max(len(r) for r in rows)
    def weight(c):  # never narrower than the longest word; long text gets more, but not proportionally
        cells = [re.sub(r"\*", "", r[c]) for r in rows if c < len(r)]
        word = max((len(w) for x in cells for w in x.split()), default=1)
        longest = max(map(len, cells), default=1)
        return word + 2.5 + 0.15 * max(0, min(longest, 40) - word)
    weights = [weight(c) for c in range(ncol)]
    widths = [int(COL_WIDTH_DXA * w / sum(weights)) for w in weights]
    widths[-1] += COL_WIDTH_DXA - sum(widths)

    tbl = copy.deepcopy(proto_tbl)
    for child in list(tbl):
        if child.tag in (qn("w:tblGrid"), qn("w:tr")):
            tbl.remove(child)
    grid = tbl.makeelement(qn("w:tblGrid"), {})
    for w in widths:
        grid.append(grid.makeelement(qn("w:gridCol"), {qn("w:w"): str(w)}))
    tbl.append(grid)
    for i, row in enumerate(rows):
        tr = tbl.makeelement(qn("w:tr"), {})
        for c in range(ncol):
            tc = tr.makeelement(qn("w:tc"), {})
            tcpr = tc.makeelement(qn("w:tcPr"), {})
            tcpr.append(tcpr.makeelement(qn("w:tcW"), {qn("w:type"): "dxa", qn("w:w"): str(widths[c])}))
            rules = (["w:top", "w:bottom"] if i == 0 else []) + (["w:bottom"] if i == len(rows) - 1 and i else [])
            if rules:
                borders = tcpr.makeelement(qn("w:tcBorders"), {})
                for side in dict.fromkeys(rules):
                    borders.append(borders.makeelement(qn(side), {qn("w:val"): "single", qn("w:sz"): "6",
                                                                  qn("w:color"): "000000"}))
                tcpr.append(borders)
            tc.append(tcpr)
            text = row[c] if c < len(row) else ""
            p = fill(cell_p, text, head_rpr if i == 0 else body_rpr)
            if i == len(rows) - 1:  # let the table break only before its last row is placed
                kn = p.find(qn("w:pPr")).find(qn("w:keepNext"))
                if kn is not None:
                    p.find(qn("w:pPr")).remove(kn)
            tc.append(p)
            tr.append(tc)
        tbl.append(tr)
    return tbl


def figure(tpl, doc, image: Path):
    holder = doc.add_paragraph()
    holder.add_run().add_picture(str(image), width=Emu(FIG_WIDTH_EMU))
    drawing_run = holder._p.find(qn("w:r"))
    p = copy.deepcopy(tpl.fig)
    for r in p.findall(qn("w:r")):
        p.remove(r)
    p.append(drawing_run)
    holder._p.getparent().remove(holder._p)
    return p


def parse(md: str):
    md = re.sub(r"<!--.*?-->", "", md, flags=re.S)
    lines = md.splitlines()
    meta, blocks, i = {}, [], 0
    while i < len(lines):
        line = lines[i].rstrip()
        i += 1
        if not line.strip():
            continue
        m = re.match(r"(TITLE|ABSTRACT|INDEX):\s*(.*)", line)
        if m:
            meta[m.group(1)] = m.group(2)
        elif line.startswith("## "):
            blocks.append(("h2", line[3:]))
        elif line.startswith("# "):
            blocks.append(("h1", line[2:]))
        elif re.match(r"TABLE [IVXL]+:", line):
            num, title = line.split(":", 1)
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):  # skip Markdown separator rows
                    rows.append(cells)
                i += 1
            blocks.append(("table", num.strip(), title.strip(), rows))
        elif line.startswith("!["):
            m = re.match(r"!\[(.*)\]\((.*)\)", line)
            blocks.append(("fig", m.group(1), m.group(2)))
        elif line.startswith("$$eq"):
            blocks.append(("eq", int(line[4:])))
        elif line.startswith("- "):
            blocks.append(("bullet", line[2:]))
        elif re.match(r"\[\d+\] ", line):
            n, text = line.split(" ", 1)
            blocks.append(("ref", n, text))
        else:
            blocks.append(("p", line))
    return meta, blocks


def build(src: Path, out: Path, kind: str, venue: str, eq_config: str):
    meta, blocks = parse(src.read_text(encoding="utf-8"))
    tpl = Template(TEMPLATE)
    doc, body = tpl.doc, tpl.body
    # title block and cover
    for p in tpl.title_paras:
        set_runs(p, [meta["TITLE"]])
    for p in tpl.kind_paras:
        set_runs(p, [kind])
    for p in tpl.venue_paras:
        set_runs(p, [venue])
    set_runs(tpl.abstract, ["Abstract", "—", meta["ABSTRACT"]])
    set_runs(tpl.index, ["Index Terms", "—", meta["INDEX"]])
    # drop the template's body, keep the final two-column section properties
    start = tpl.els.index(tpl.first_body)
    for el in tpl.els[start:-1]:
        body.remove(el)
    anchor = tpl.final_sect

    def add(el):
        anchor.addprevious(el)

    for b in blocks:
        kind_ = b[0]
        if kind_ == "h1":
            add(fill(tpl.h1, b[1]))
        elif kind_ == "h2":
            add(fill(tpl.h2, b[1]))
        elif kind_ == "p":
            add(fill(tpl.body_p, b[1]))
        elif kind_ == "bullet":
            m = re.match(r"\*([^*]+)\*\s*(.*)", b[1])
            lead, rest = (f"• {m.group(1)} ", m.group(2)) if m else ("• ", b[1])
            p = fill(tpl.bullet, rest)  # plain text, *italic* spans
            first = p.find(qn("w:r"))
            make_run(p, lead, first_rpr(tpl.bullet), italic=bool(m))
            if first is not None:  # make_run appends; move the lead to the front
                first.addprevious(p.findall(qn("w:r"))[-1])
            add(p)
        elif kind_ == "table":
            add(fill(tpl.tab_num, b[1]))
            add(fill(tpl.tab_title, b[2]))
            add(table(tpl, b[3]))
        elif kind_ == "fig":
            add(figure(tpl, doc, (src.parent / b[2]).resolve()))
            add(fill(tpl.fig_cap, b[1]))
        elif kind_ == "eq":
            p = copy.deepcopy(tpl.eq[b[1] - 1])
            for t in p.iter(qn("w:t")):
                t.text = t.text.replace("D1+D2", eq_config)
            add(p)
        elif kind_ == "ref":
            add(fill(tpl.ref, f"{b[1]}\t{b[2]}"))
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    counts = {k: sum(1 for b in blocks if b[0] == k) for k in ("h1", "p", "table", "fig", "ref")}
    print(f"wrote {out}: {counts}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("src", type=Path)
    ap.add_argument("out", type=Path)
    ap.add_argument("--kind", default="Final Report", help="cover-page line, e.g. 'Final Report'")
    ap.add_argument("--venue", default="ESP3201A Final Report, National University of Singapore")
    ap.add_argument("--eq-config", default="D1h+D2", help="configuration named in equations (1) and (2)")
    a = ap.parse_args()
    build(a.src, a.out, a.kind, a.venue, a.eq_config)
